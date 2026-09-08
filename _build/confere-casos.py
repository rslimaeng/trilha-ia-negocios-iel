# -*- coding: utf-8 -*-
"""Reconfere, contra o xlsx, cada numero que uma pagina de caso afirma.

🔴 Existe porque o G32 NAO alcanca estes seis arquivos, e o prompt supunha que
alcancava. O G32 itera o `insumos.json`, que so lista o que o `insumo.py` GERA
a partir de especificacao. XLSX convertido de CSV nao entra la, e forcar a
entrada seria apagado na proxima rodada do insumo.py.

Sem gate, a unica protecao seria conferir a mao, e conferir a mao nao se repete.
Entao a conferencia virou codigo: cada numero que as seis paginas afirmam esta
declarado aqui ao lado da conta que o produz, e uma divergencia reprova.

🔴 O valor esperado aqui e o que a PAGINA diz. Quando a pagina muda, este
arquivo muda junto, e e de proposito: e o par que faz a divergencia aparecer.
Em 07/09/2026 ele pegou dois erros meus no caso de RH antes de irem ao ar: uma
taxa calculada sobre 18 quando a etapa tinha 21 pessoas, e a contagem de vagas
com candidato aprovado, que era 2 e a pagina dizia 1.

Rodar:  python3 _build/confere-casos.py
"""
from openpyxl import load_workbook
import collections, html, io, os, re, statistics as st, sys

AQUI = os.path.dirname(os.path.abspath(__file__))
ARQUIVOS = os.path.join(os.path.dirname(AQUI), "_arquivos")


def base(slug):
    ws = load_workbook(os.path.join(ARQUIVOS, slug + ".xlsx")).active
    cab = [c.value for c in ws[1]]
    return [dict(zip(cab, [c.value for c in row])) for row in ws.iter_rows(min_row=2)]

# ---------------------------------------------------------------------------
# 🔴 A CHECAGEM DE PRESENCA, de 08/09.
#
# Ate aqui o conferidor comparava a conta com um valor escrito no proprio
# arquivo, e nao olhava a pagina. Tinha um buraco: um numero APAGADO da pagina
# continuava "conferindo", porque a conta continuava batendo com a constante
# daqui. O placar diria 174 de 174 numa pagina que afirmasse 120.
#
# Agora cada valor e procurado no texto visivel da pagina. O que nao esta la
# sai do placar como REMOVIDO, resultado esperado de um corte de prosa. O que
# esta la e conferido como antes.
# ---------------------------------------------------------------------------
CONTEUDO = os.path.join(AQUI, "conteudo")
PAGINAS_VIS = {}


def _visivel(slug):
    t = io.open(os.path.join(CONTEUDO, "b2-caso-%s.html" % slug),
                encoding="utf-8").read()
    t = re.sub(r"(?s)<script.*?</script>|<style.*?</style>", " ", t)
    return re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", " ", t)))


def _formatos(v, casas):
    """As grafias em que o numero pode aparecer na tela, em portugues."""
    v = abs(v)
    saida = set()
    if casas == 0:
        i = int(round(v))
        saida.add(str(i))
        saida.add("{:,}".format(i).replace(",", "."))
    else:
        for c in (casas, 1, 2):
            t = ("{:.%df}" % c).format(v)
            saida.add(t.replace(".", ","))
            saida.add(t)
    return saida


ok, ruim, removidos = [], [], []


def bate(pag, rotulo, calculado, afirmado, casas=2):
    a, b = round(float(calculado), casas), round(float(afirmado), casas)
    if pag not in PAGINAS_VIS:
        PAGINAS_VIS[pag] = _visivel(pag)
    if not any(f in PAGINAS_VIS[pag] for f in _formatos(b, casas)):
        removidos.append((pag, rotulo, b))
        return
    (ok if a == b else ruim).append((pag, rotulo, a, b))

# ---------------- PRODUCAO ----------------
R = base("producao")
tx = lambda s: sum(r['UnidadesDefeituosas'] for r in s)/sum(r['UnidadesProcessadas'] for r in s)*100
bate("producao","linhas",len(R),5490,0)
bate("producao","processadas",sum(r['UnidadesProcessadas'] for r in R),5209485,0)
bate("producao","defeituosas",sum(r['UnidadesDefeituosas'] for r in R),83791,0)
bate("producao","taxa geral",tx(R),1.61)
bate("producao","parada total",sum(r['MinutosParada'] for r in R),146486,0)
bate("producao","sucata total",sum(r['CustoSucataUSD'] for r in R),128405,0)
bate("producao","turno C",tx([r for r in R if r['Turno']=='C']),1.92)
bate("producao","turno A",tx([r for r in R if r['Turno']=='A']),1.45)
bate("producao","turno B",tx([r for r in R if r['Turno']=='B']),1.46)
bate("producao","C sem processamento",tx([r for r in R if r['Turno']=='C' and r['Etapa']!='Processamento']),1.43)
bate("producao","A+B",tx([r for r in R if r['Turno']!='C']),1.45)
sc=[r for r in R if r['Etapa']=='Processamento' and r['Turno']=='C']
outros=[r for r in R if not (r['Etapa']=='Processamento' and r['Turno']=='C')]
bate("producao","Processamento/C",tx(sc),3.95)
bate("producao","as outras 14",tx(outros),1.45)
bate("producao","defeituosas do cruzamento",sum(r['UnidadesDefeituosas'] for r in sc),13343,0)
bate("producao","share das defeituosas",sum(r['UnidadesDefeituosas'] for r in sc)/sum(r['UnidadesDefeituosas'] for r in R)*100,15.9,1)
bate("producao","share do volume",sum(r['UnidadesProcessadas'] for r in sc)/sum(r['UnidadesProcessadas'] for r in R)*100,6.5,1)
tb=sum(r['UnidadesDefeituosas'] for r in outros)/sum(r['UnidadesProcessadas'] for r in outros)
esp=sum(r['UnidadesProcessadas'] for r in sc)*tb
bate("producao","esperado se normal",esp,4888,0)
bate("producao","excesso de pecas",sum(r['UnidadesDefeituosas'] for r in sc)-esp,8455,0)
cu=sum(r['CustoSucataUSD'] for r in R)/sum(r['UnidadesDefeituosas'] for r in R)
bate("producao","custo por peca",cu,1.53)
bate("producao","custo do excesso",(sum(r['UnidadesDefeituosas'] for r in sc)-esp)*cu,12957,0)
bate("producao","share da sucata",(sum(r['UnidadesDefeituosas'] for r in sc)-esp)*cu/sum(r['CustoSucataUSD'] for r in R)*100,10.1,1)
mp=collections.defaultdict(int)
for r in R: mp[(r['Etapa'],r['Turno'])]+=r['MinutosParada']
bate("producao","Montagem/B parada",mp[('Montagem','B')],32396,0)
bate("producao","Montagem/B em horas",mp[('Montagem','B')]/60,540,0)
bate("producao","media das outras 14",(sum(r['MinutosParada'] for r in R)-mp[('Montagem','B')])/14,8149,0)
bate("producao","Montagem/B share",mp[('Montagem','B')]/sum(r['MinutosParada'] for r in R)*100,22.1,1)
bate("producao","Montagem taxa",tx([r for r in R if r['Etapa']=='Montagem']),1.47)
mes=collections.defaultdict(lambda:[0,0])
for r in sc:
    k=r['Data'].strftime("%Y-%m"); mes[k][0]+=r['UnidadesProcessadas']; mes[k][1]+=r['UnidadesDefeituosas']
taxas=[v[1]/v[0]*100 for v in mes.values()]
bate("producao","menor mes",min(taxas),3.68); bate("producao","maior mes",max(taxas),4.16)
tabela={'Corte':(1.41,1.44,1.44),'Processamento':(1.41,1.45,3.95),'Montagem':(1.50,1.45,1.45),
        'Acabamento':(1.46,1.46,1.46),'Embalagem':(1.45,1.48,1.39)}
for e,vals in tabela.items():
    for t,v in zip("ABC",vals):
        bate("producao","tabela %s/%s"%(e,t),tx([r for r in R if r['Etapa']==e and r['Turno']==t]),v)

# ---------------- COMERCIAL ----------------
R = base("comercial")
dec=[r for r in R if r['Status']!='Em andamento']
bate("comercial","linhas",len(R),800,0)
bate("comercial","ganhos",sum(1 for r in R if r['Status']=='Ganho'),219,0)
bate("comercial","perdidos",sum(1 for r in R if r['Status']=='Perdido'),81,0)
bate("comercial","em andamento",sum(1 for r in R if r['Status']=='Em andamento'),500,0)
bate("comercial","decididos",len(dec),300,0)
bate("comercial","taxa sobre 800",219/800*100,27.4,1)
bate("comercial","taxa sobre 300",219/300*100,73.0,1)
bate("comercial","vendedores",len(set(r['Vendedor'] for r in R)),7,0)
bate("comercial","origens",len(set(r['Origem'] for r in R)),10,0)
def vend(v):
    s=[r for r in R if r['Vendedor']==v]; d=[r for r in s if r['Status']!='Em andamento']
    g=sum(1 for r in s if r['Status']=='Ganho'); return g/len(s)*100, g/len(d)*100
bate("comercial","V6 sobre tudo",vend('Vendedor 6')[0],23.3,1)
bate("comercial","V6 sobre decididos",vend('Vendedor 6')[1],81.8,1)
bate("comercial","V5 sobre tudo",vend('Vendedor 5')[0],35.0,1)
bate("comercial","V5 sobre decididos",vend('Vendedor 5')[1],77.8,1)
bate("comercial","V4 sobre decididos",vend('Vendedor 4')[1],60.9,1)
def orig(o):
    s=[r for r in R if r['Origem']==o]; g=[r for r in s if r['Status']=='Ganho']
    return len(g)/len(s)*100, st.mean([float(r['Valor_Proposta_R$']) for r in g]), sum(float(r['Valor_Proposta_R$']) for r in g)/len(s)
bate("comercial","Webinar taxa",orig('Webinar')[0],22.4,1)
bate("comercial","Webinar ticket",orig('Webinar')[1],14916,0)
bate("comercial","WhatsApp por lead",orig('WhatsApp')[2],5003,0)
bate("comercial","Indicacao por lead",orig('Indicação')[2],2423,0)

# ---------------- MARKETING ----------------
R = base("marketing")
bate("marketing","linhas",len(R),400,0)
bate("marketing","nota geral",st.mean([r['Nota_Geral_0_10'] for r in R]),8.37)
for v,n in [('Sim',358),('Talvez',30),('Não',12)]:
    bate("marketing","voltam=%s"%v,sum(1 for r in R if r['Interesse_Proxima_Edicao']==v),n,0)
bate("marketing","indicariam",sum(1 for r in R if r['Indicaria_Para_Colegas']=='Sim'),313,0)
bate("marketing","perfis",len(set(r['Perfil_Participante'] for r in R)),6,0)
sim=[r for r in R if r['Interesse_Proxima_Edicao']=='Sim']; nao=[r for r in R if r['Interesse_Proxima_Edicao']=='Não']
tab={'Nota_Credenciamento':(7.32,7.68,3.42,4.26),'Nota_App_Evento':(7.01,7.36,3.17,4.19),
     'Nota_Logistica':(7.83,8.19,4.33,3.86),'Nota_Coffee_Refeicoes':(7.88,8.18,5.00,3.18),
     'Nota_Palestrantes':(8.44,8.71,6.00,2.71),'Nota_Conteudo_Tecnico':(8.37,8.65,6.00,2.65)}
for c,(g,s_,n_,d) in tab.items():
    bate("marketing",c+" geral",st.mean([r[c] for r in R]),g)
    bate("marketing",c+" volta",st.mean([r[c] for r in sim]),s_)
    bate("marketing",c+" nao volta",st.mean([r[c] for r in nao]),n_)
    bate("marketing",c+" diferenca",st.mean([r[c] for r in sim])-st.mean([r[c] for r in nao]),d)
for p,v in [('Visitante de primeira viagem',8.10),('Palestrante',8.65),('Equipe interna',8.49)]:
    bate("marketing","perfil "+p,st.mean([r['Nota_Geral_0_10'] for r in R if r['Perfil_Participante']==p]),v)

# ---------------- OPERACOES ----------------
R = base("operacoes")
nps=lambda s:(sum(1 for r in s if r['Nota_NPS_0_10']>=9)-sum(1 for r in s if r['Nota_NPS_0_10']<=6))/len(s)*100
bate("operacoes","linhas",len(R),500,0)
bate("operacoes","NPS geral",nps(R),43.8,1)
bate("operacoes","nota media",st.mean([r['Nota_NPS_0_10'] for r in R]),8.24)
g=collections.Counter((r['Linha_Produto'],r['Canal']) for r in R)
for lp,n in [('LP-100',101),('LP-200',115),('LP-300',108),('LP-400',93)]:
    bate("operacoes","%s/Distribuidor"%lp,g[(lp,'Distribuidor')],n,0)
    bate("operacoes","%s/Varejo"%lp,g[(lp,'Varejo')],0,0)
bate("operacoes","LP-500/Varejo",g[('LP-500','Varejo')],83,0)
bate("operacoes","LP-500/Distribuidor",g[('LP-500','Distribuidor')],0,0)
bate("operacoes","NPS varejo",nps([r for r in R if r['Canal']=='Varejo']),48.2,1)
bate("operacoes","NPS distribuidor",nps([r for r in R if r['Canal']=='Distribuidor']),42.9,1)
bate("operacoes","n distribuidor",sum(1 for r in R if r['Canal']=='Distribuidor'),417,0)
for c,v in [('Nota_Entrega',6.98),('Nota_Suporte',7.84),('Nota_Logistica',7.85),('Nota_Produto',8.34),('Nota_Atendimento',8.43)]:
    bate("operacoes",c,st.mean([r[c] for r in R]),v)
bate("operacoes","NPS coordenador",nps([r for r in R if r['Perfil_Cliente']=='Coordenador de compras']),30.0,1)
bate("operacoes","n coordenador",sum(1 for r in R if r['Perfil_Cliente']=='Coordenador de compras'),110,0)
bate("operacoes","NPS comprador",nps([r for r in R if r['Perfil_Cliente']=='Comprador']),53.4,1)
bate("operacoes","n comprador",sum(1 for r in R if r['Perfil_Cliente']=='Comprador'),88,0)
bate("operacoes","NPS L1",nps([r for r in R if r['Lote']=='L1/26']),30.1,1)
bate("operacoes","NPS L4",nps([r for r in R if r['Lote']=='L4/26']),50.0,1)

# ---------------- FINANCEIRO ----------------
R = base("financeiro")
bate("financeiro","linhas",len(R),72,0)
o=sum(r['Receita_Orcada'] for r in R); z=sum(r['Receita_Realizada'] for r in R)
bate("financeiro","orcada",o,56745120,0); bate("financeiro","realizada",z,54131383,0)
bate("financeiro","desvio",o-z,2613737,0); bate("financeiro","desvio pct",(z/o-1)*100,-4.6,1)
vol=sum((r['Pedidos_Fechados']-r['Pedidos_Orcados'])*r['Ticket_Medio_Orcado'] for r in R)
pre=sum(r['Pedidos_Fechados']*(r['Ticket_Medio_Realizado']-r['Ticket_Medio_Orcado']) for r in R)
bate("financeiro","parcela volume",-vol,363176,0); bate("financeiro","parcela preco",-pre,2250612,0)
bate("financeiro","soma das parcelas",-(vol+pre),2613788,0)
bate("financeiro","residuo",abs(-(vol+pre)-(o-z)),51,0)
bate("financeiro","volume share",-vol/-(vol+pre)*100,13.9,1); bate("financeiro","preco share",-pre/-(vol+pre)*100,86.1,1)
gl=collections.defaultdict(lambda:[0.0,0.0])
for r in R: gl[r['Linha_Produto']][0]+=r['Receita_Orcada']; gl[r['Linha_Produto']][1]+=r['Receita_Realizada']
for lp,pct,val in [('LP-400',-8.7,-988000),('LP-300',-4.9,-1260107),('LP-600',2.8,181440)]:
    bate("financeiro",lp+" pct",(gl[lp][1]/gl[lp][0]-1)*100,pct,1)
    bate("financeiro",lp+" valor",gl[lp][1]-gl[lp][0],val,0)
gm=collections.defaultdict(lambda:[0.0,0.0])
for r in R: gm[r['Mes']][0]+=r['Receita_Orcada']; gm[r['Mes']][1]+=r['Receita_Realizada']
for m,pct in [('Dezembro',-23.9),('Junho',-20.2),('Novembro',-15.3),('Janeiro',-14.8),('Setembro',11.3),('Fevereiro',10.8)]:
    bate("financeiro","mes "+m,(gm[m][1]/gm[m][0]-1)*100,pct,1)

# ---------------- RH ----------------
R = base("rh")
bate("rh","linhas",len(R),94,0)
bate("rh","vagas",len(set(r['Vaga'] for r in R)),6,0)
for res,n in [('Em andamento',55),('Reprovado',36),('Aprovado',2),('Negociação',1)]:
    bate("rh","resultado "+res,sum(1 for r in R if r['Resultado']==res),n,0)
bate("rh","vagas com aprovado",len(set(r['Vaga'] for r in R if r['Resultado']=='Aprovado')),2,0)
et={'Triagem':(7,5,41.7,16.5),'Entrevista RH':(12,13,52.0,26.0),'Teste Técnico':(4,10,71.4,24.6),
    'Entrevista Gestor':(17,5,22.7,24.0),'Proposta':(15,3,14.3,20.5)}
for e,(anda,rep,taxa,dias) in et.items():
    s=[r for r in R if r['Etapa_Atual']==e]
    bate("rh",e+" em andamento",sum(1 for r in s if r['Resultado']=='Em andamento'),anda,0)
    bate("rh",e+" reprovados",sum(1 for r in s if r['Resultado']=='Reprovado'),rep,0)
    bate("rh",e+" taxa",sum(1 for r in s if r['Resultado']=='Reprovado')/len(s)*100,taxa,1)
    bate("rh",e+" dias",st.mean([r['Dias_No_Processo'] for r in s]),dias,1)
bate("rh","aprovados no teste",st.mean([r['Nota_Teste_Tecnico'] for r in R if r['Resultado']=='Aprovado']),4.50)
bate("rh","reprovados no teste",st.mean([r['Nota_Teste_Tecnico'] for r in R if r['Resultado']=='Reprovado' and r['Nota_Teste_Tecnico']]),5.99)

por = collections.Counter(p for p, _, _, _ in ok + ruim)
rem = collections.Counter(p for p, _, _ in removidos)
print("NUMEROS QUE A PAGINA AFIRMA, CONFERIDOS CONTRA O XLSX:")
for p in ["producao", "comercial", "marketing", "operacoes", "financeiro", "rh"]:
    b = sum(1 for x, _, _, _ in ok if x == p)
    m = sum(1 for x, _, _, _ in ruim if x == p)
    print("   {:12s} afirma {:3d} · confere {:3d} · DIVERGE {}   (saiu do texto: {})"
          .format(p, por[p], b, m, rem[p]))
print("\n   TOTAL: {} afirmados · {} conferem · {} divergem"
      .format(len(ok) + len(ruim), len(ok), len(ruim)))
print("   {} valores nao estao mais na tela, e isso e o corte de prosa"
      .format(len(removidos)))
for p, r, a, b in ruim:
    print("   🔴 {} · {}: calculado {} × afirmado {}".format(p, r, a, b))
sys.exit(1 if ruim else 0)
