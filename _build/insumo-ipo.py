# -*- coding: utf-8 -*-
"""Gera o `ipo-do-meu-processo.xlsx`, o instrumento da aula 3 do B3.

🔴 NAO E INSUMO DE ANALISE. As outras planilhas do site trazem dados sujos de
proposito, para a pessoa achar armadilha. Esta e o contrario: e uma FOLHA EM
BRANCO com um exemplo ao lado, e a sujeira aqui seria defeito.

Duas abas, e a ordem importa: o exemplo vem primeiro porque a folha em branco
sozinha e a "tela em branco" que a propria aula 1 ensinou a reconhecer.

A pessoa carrega este arquivo por tres aulas: preenche na 3, mede na 4 e
precifica na 5. E a mesma tabela que o bloco 4 vira instrucao de assistente.
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SAIDA = os.path.join(RAIZ, "_arquivos", "ipo-do-meu-processo.xlsx")

CABECA = ["Campo", "O que eu escrevi"]
CAMPOS = ["Área", "Entrada", "Saída", "Exceção", "Recursos"]

PASSOS_CAB = ["#", "O passo (verbo de ação + critério)", "IN/EX", "Tempo", "É medido ou é chute?"]

EXEMPLO_CAMPOS = [
    ("Área", "Comercial"),
    ("Entrada", "Pedido de orçamento do cliente, por e-mail ou WhatsApp, às vezes com "
                "foto de uma lista escrita à mão"),
    ("Saída", "Proposta em PDF, com itens, valor total, prazo de entrega e validade, "
              "enviada ao cliente"),
    ("Exceção", "Item fora da tabela vigente vira consulta ao técnico. Desconto acima de "
                "10% precisa do gerente. Cliente novo exige análise de crédito antes de enviar"),
    ("Recursos", "Tabela de preços vigente, modelo de proposta, acesso ao histórico do cliente"),
]

EXEMPLO_PASSOS = [
    (1, "Ler o pedido e listar os itens; se vier como foto, transcrever a lista antes", "IN", "15 min", "medido"),
    (2, "Conferir cada item contra a tabela vigente; marcar os que não estão nela", "IN", "20 min", "medido"),
    (3, "Consultar preço e prazo com o técnico da linha, para os itens marcados", "EX", "2 dias", "medido"),
    (4, "Montar a proposta no modelo, com validade de 15 dias", "IN", "40 min", "chute"),
    (5, "Revisar valor e prazo contra o pedido original, e enviar", "IN", "10 min", "medido"),
]

AZUL = "164194"
CINZA = "F4F6FA"


def _titulo(ws, linha, texto):
    c = ws.cell(row=linha, column=1, value=texto)
    c.font = Font(bold=True, size=12, color=AZUL)


def _cabecalho(ws, linha, colunas):
    for i, nome in enumerate(colunas, start=1):
        c = ws.cell(row=linha, column=i, value=nome)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(vertical="center", wrap_text=True)


def _aba(wb, nome, campos, passos):
    ws = wb.create_sheet(nome)
    ws.sheet_view.showGridLines = False

    _titulo(ws, 1, "1 · As quatro pontas do processo")
    _cabecalho(ws, 2, CABECA)
    for i, (campo, valor) in enumerate(campos, start=3):
        ws.cell(row=i, column=1, value=campo).font = Font(bold=True)
        c = ws.cell(row=i, column=2, value=valor)
        c.alignment = Alignment(wrap_text=True, vertical="top")

    inicio = 3 + len(campos) + 1
    _titulo(ws, inicio, "2 · Os passos, de 4 a 7. Se passar de 7, são dois processos")
    _cabecalho(ws, inicio + 1, PASSOS_CAB)
    for i, linha in enumerate(passos, start=inicio + 2):
        for j, valor in enumerate(linha, start=1):
            c = ws.cell(row=i, column=j, value=valor)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if j == 1:
                c.alignment = Alignment(horizontal="center", vertical="top")

    for larg, col in zip([13, 78, 8, 12, 20], range(1, 6)):
        ws.column_dimensions[get_column_letter(col)].width = larg
    return ws


def grava():
    wb = Workbook()
    wb.remove(wb.active)

    _aba(wb, "Exemplo · orçamento", EXEMPLO_CAMPOS, EXEMPLO_PASSOS)
    # a folha em branco leva 7 linhas de passo: o teto da regra, nao a media,
    # para ninguem achar que 5 e o limite.
    vazios = [(n, None, None, None, None) for n in range(1, 8)]
    ws = _aba(wb, "O meu processo", [(c, None) for c in CAMPOS], vazios)
    ws.cell(row=1, column=4, value="Dados fictícios no exemplo. Esta aba é sua.").font = \
        Font(italic=True, size=9, color="6B6B6B")

    os.makedirs(os.path.dirname(SAIDA), exist_ok=True)
    wb.save(SAIDA)
    print("gravado: %s" % os.path.relpath(SAIDA, RAIZ))
    for aba in wb.sheetnames:
        print("  aba %-24s %d linhas" % (aba, wb[aba].max_row))


if __name__ == "__main__":
    grava()
